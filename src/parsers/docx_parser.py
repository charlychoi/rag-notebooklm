"""DOCX 파서 — python-docx 기반."""
from __future__ import annotations

from pathlib import Path

from src.parsers.base import ParsedDocument, ParserUnavailableError


class DocxParser:
    def __init__(self) -> None:
        try:
            import docx as _docx  # noqa: F401
        except ImportError:
            raise ParserUnavailableError("python-docx 미설치: pip install python-docx")

    def parse(self, path: Path) -> ParsedDocument:
        import docx
        doc = docx.Document(str(path))
        lines: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            if "Heading 1" in style:
                lines.append(f"# {text}")
            elif "Heading 2" in style:
                lines.append(f"## {text}")
            elif "Heading 3" in style:
                lines.append(f"### {text}")
            else:
                lines.append(text)
        markdown = "\n\n".join(lines)
        return ParsedDocument(
            source_id=path.stem,
            title=path.stem.replace("_", " "),
            markdown=markdown,
            metadata={"source": path.name, "format": "docx"},
        )
